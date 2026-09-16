from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "QA" / "output" / "BOUESTI-ExamPilot-Timetable-Documentation.docx"
PREVIOUS = Path("C:/Users/burinious/Downloads/WhatsApp Image 2026-08-23 at 06.57.51 (2).jpeg")
CURRENT = ROOT / "QA" / "output" / "QA-BOUESTI-timetable-page-1.png"
RECON = ROOT / "QA" / "output" / "BOUESTI-timetable-reconciliation.xlsx"

NAVY = "12253F"
TEAL = "2F7E8A"
LIGHT = "E8EDF2"
PALE = "F4F7F9"
GRAY = "64748B"
BORDER = "D9D9D9"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def style_run(run, size=10, bold=False, color="000000", italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def add_text(p, text, size=10, bold=False, color="000000", italic=False):
    run = p.add_run(text)
    style_run(run, size, bold, color, italic)
    return run

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, 16 if level == 1 else 11, True, NAVY if level == 1 else TEAL)
    return p

def body(doc, text, size=10, color="26364A"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text, size, False, color)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text, 9.5, False, "26364A")
    return p

def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, text, 8.5, False, GRAY, True)
    return p

def picture(doc, path, width, description):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    docPr = shape._inline.docPr
    docPr.set("descr", description)
    return p

def table(doc, headers, rows, widths=None, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        if widths: cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, label, font_size, True, "FFFFFF")
    for r_idx, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths: cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, PALE if r_idx % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, str(value), font_size, False, "26364A")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_meta_table(doc):
    rows = [
        ("Previous timetable", "2024/2025 | Second Semester | 6–7 October 2025"),
        ("Current ExamPilot cycle", "2026/2027 | First Semester | 1–2 February 2027"),
        ("Current release label", "Revision 1 | DRAFT | Unpublished"),
        ("Source decision", "Retain current approved records; use the photograph as the layout and operations reference."),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_repeat_table_header(t.rows[0])
    for i, label in enumerate(("Field", "Value")):
        cell = t.rows[0].cells[i]
        cell.width = Inches(1.75 if i == 0 else 5.05)
        set_cell_border(cell)
        set_cell_margins(cell, 110, 130, 110, 130)
        set_cell_shading(cell, NAVY)
        add_text(cell.paragraphs[0], label, 9, True, "FFFFFF")
    for i, (key, value) in enumerate(rows):
        cells = t.add_row().cells
        for j, cell in enumerate(cells):
            cell.width = Inches(1.75 if j == 0 else 5.05)
            set_cell_border(cell)
            set_cell_margins(cell, 110, 130, 110, 130)
            set_cell_shading(cell, LIGHT if j == 0 else (PALE if i % 2 else "FFFFFF"))
            p = cell.paragraphs[0]
            add_text(p, key if j == 0 else value, 9, j == 0, NAVY if j == 0 else "26364A")
    return t

def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.color.rgb = RGBColor.from_string("000000")

def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "BOUESTI EXAMPILOT  |  TIMETABLE QA DOCUMENTATION", 7.5, True, GRAY)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    add_text(p, "Internal review document  |  Page ", 8, False, GRAY)
    add_page_number(p)

def build():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    add_header_footer(section)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "BOUESTI ExamPilot Timetable Documentation", 28, True, "000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    add_text(p, "Print workflow review and QA evidence", 14, False, TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "Bamidele Olumilua University of Education, Science and Technology", 10, True, NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    add_text(p, "Ikere-Ekiti  |  College of Science  |  14 September 2026", 9, False, GRAY)
    body(doc, "This document records the print-workflow update made after comparing the previous BOUESTI timetable photograph with the current ExamPilot output. It identifies the examination cycle represented by each source, documents the decision to preserve the current cycle, and provides visual and structural evidence for the regenerated exports.", 10.5)
    add_meta_table(doc)
    body(doc, "Review outcome: the current ExamPilot records remain in place. The photograph is used to guide layout, session grouping, split-venue presentation, and operational staffing fields. Replacing the current course or date data is blocked until approved 2024/2025 source records are supplied.", 10.5)

    doc.add_page_break()
    heading(doc, "Cycle Decision and Scope")
    body(doc, "The two supplied documents describe separate examination cycles. The photograph is headed 2024/2025 Second Semester and shows 6–7 October 2025. The ExamPilot PDF is headed 2026/2027 First Semester Examination and contains dates beginning 1 February 2027. The current records therefore remain the source for the updated print workflow.")
    table(doc, ["Dimension", "Previous timetable photograph", "Current ExamPilot printout", "Decision"], [
        ("Academic cycle", "2024/2025, Second Semester", "2026/2027, First Semester", "Retain current cycle"),
        ("Dates shown", "6–7 October 2025", "1–2 February 2027 in scheduled rows", "Do not replace"),
        ("Session windows", "08:30–11:30; 12:00–14:00; 14:30–17:30", "09:00–12:00; 14:00–17:00 in approved slots", "Print approved slots; flag missing midday data"),
        ("Course coverage", "Multiple 100–500 level courses", "CSC 202, CSC 405, MTH 202 scheduled; PHY 202 remains unscheduled", "Do not copy photo courses"),
    ], [1.1, 1.8, 2.25, 1.65], 8.1)
    heading(doc, "Scope Controls", 2)
    bullet(doc, "No course registrations, programme assignments, venue allocations, or invigilator assignments were invented.")
    bullet(doc, "The export allowlist contains operational timetable fields only and does not expose student names or matriculation numbers.")
    bullet(doc, "The release status remains configurable and visibly truthful: Revision 1, DRAFT, Unpublished.")

    doc.add_page_break()
    heading(doc, "Visual Evidence")
    body(doc, "The paired screenshots show the visual baseline and the regenerated ExamPilot output. The photograph is retained for the operational structure and typography cues. The regenerated page applies the same practical information hierarchy while making the current cycle explicit.")
    if PREVIOUS.exists() and CURRENT.exists():
        t = doc.add_table(rows=2, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_repeat_table_header(t.rows[0])
        for i, label in enumerate(("Previous timetable photograph", "Regenerated ExamPilot output")):
            cell = t.rows[0].cells[i]
            cell.width = Inches(3.25)
            set_cell_shading(cell, NAVY)
            set_cell_border(cell)
            set_cell_margins(cell, 70, 70, 70, 70)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, label, 8, True, "FFFFFF")
        for cell in t.rows[1].cells:
            cell.width = Inches(3.25)
            set_cell_border(cell, "FFFFFF", "0")
            set_cell_margins(cell, 30, 30, 30, 30)
        p1 = t.cell(1, 0).paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER; r1 = p1.add_run(); shape1 = r1.add_picture(str(PREVIOUS), width=Inches(3.0)); shape1._inline.docPr.set("descr", "Previous BOUESTI timetable photograph")
        p2 = t.cell(1, 1).paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; r2 = p2.add_run(); shape2 = r2.add_picture(str(CURRENT), width=Inches(3.0)); shape2._inline.docPr.set("descr", "Regenerated ExamPilot timetable PDF screenshot")
    caption(doc, "Figure 1. Previous BOUESTI timetable structure beside the regenerated ExamPilot printout.")
    heading(doc, "Observed Visual Changes", 2)
    table(doc, ["Element", "Regenerated behavior"], [
        ("Institutional identity", "Full university name, Ikere-Ekiti, and College of Science appear above the timetable."),
        ("Cycle and release state", "Session, semester, examination period, revision, status, and publication state are visible together."),
        ("Session grouping", "Each approved session window is printed as a separate band. Empty approved windows remain visible when present."),
        ("Split venues", "CSC 202 is one examination group with LLT1 300/415 and Navates 2 115/415 allocations."),
        ("Staffing", "Chief and venue-specific staffing fields use existing assignments; absent chief coverage prints Unassigned."),
        ("Footer", "Page numbering is generated from the final PDF page count, avoiding the previous footer-only page."),
    ], [1.65, 5.15], 8.7)

    doc.add_page_break()
    heading(doc, "Current Operational View")
    body(doc, "The regenerated printout contains four scheduled venue rows from the selected local generation. The aggregate event dataset contains 1,335 candidates across four events; 715 candidates are allocated in the scheduled rows shown below. PHY 202 remains unscheduled under the current approved constraints and is not silently printed as assigned.")
    table(doc, ["Date", "Window", "Course", "Venue allocation", "Candidates", "Staffing"], [
        ("2027-02-01", "09:00–12:00", "CSC 202 | Data Structures", "LLT1", "300 of 415", "Chief Unassigned | STAFF-001"),
        ("2027-02-01", "09:00–12:00", "CSC 202 | Data Structures", "Navates 2", "115 of 415", "Chief Unassigned | STAFF-002"),
        ("2027-02-01", "14:00–17:00", "CSC 405 | Distributed Systems", "S1", "60 of 60", "Chief Unassigned | STAFF-001"),
        ("2027-02-02", "09:00–12:00", "MTH 202 | Linear Algebra", "LLT1", "240 of 240", "Chief Unassigned | STAFF-001"),
    ], [0.82, 0.9, 1.85, 1.0, 0.9, 1.23], 8.1)
    heading(doc, "Data Blockers", 2)
    table(doc, ["Severity", "Blocker", "Evidence", "Required action"], [
        ("High", "Chief invigilator coverage is missing for four exported rows.", "Rows show Chief = Unassigned; venue codes remain STAFF-001 or STAFF-002 where linked.", "Load approved chief assignments before operational release."),
        ("Medium", "The current approved snapshot has no 12:00–14:00 slot.", "Session Windows contains only 09:00–12:00 and 14:00–17:00.", "Add the approved midday slot if this cycle requires it."),
        ("Medium", "The photo and PDF are different cycles.", "2024/2025 versus 2026/2027 headings and dates.", "Supply approved 2024/2025 records only if reproduction is intended."),
    ], [0.62, 2.05, 2.35, 1.68], 8.0)

    doc.add_page_break()
    heading(doc, "Verification Results")
    table(doc, ["Check", "Result", "Evidence"], [
        ("Typecheck", "Passed", "npm run typecheck"),
        ("Lint", "Passed", "npm run lint"),
        ("Focused timetable/export tests", "Passed", "19 tests across export-publication, timetable-review, and aggregate-timetable"),
        ("PDF structural check", "Passed", "PDF 1.3; header %PDF; EOF marker present; current output is 1 page"),
        ("PDF visual check", "Passed", "Rendered page inspected at full resolution; no orphan footer; headings and staffing columns visible"),
        ("Excel inspection", "Passed", "Six sheets; Full Timetable has 8 rows and 14 columns"),
        ("CSV inspection", "Passed", "Five lines including header; 14 operational columns; privacy scan clear"),
        ("Split-venue totals", "Passed", "CSC 202 total 415 equals LLT1 300 plus Navates 2 115"),
        ("Privacy check", "Passed", "No student names, matric numbers, or registration fields in exports"),
        ("git diff --check", "Passed", "No whitespace errors; existing LF/CRLF warnings only"),
    ], [2.0, 0.8, 4.1], 8.2)
    heading(doc, "Export Set", 2)
    table(doc, ["Artifact", "Purpose"], [
        ("QA-BOUESTI-timetable.pdf", "Operational printout with institutional heading, session bands, split allocations, and staffing fields."),
        ("QA-BOUESTI-timetable.xlsx", "Workbook with full, programme, venue, CBT, session-window, and diagnostic views."),
        ("QA-BOUESTI-timetable-export.csv", "Privacy-safe native CSV export aligned to the Excel schema."),
        ("BOUESTI-timetable-reconciliation.xlsx / .csv", "Comparison of the previous photograph and current cycle values with source-of-truth statuses."),
        ("BOUESTI-timetable-verification.md", "Issue log with severity, evidence, fix state, and approved-data requirements."),
    ], [2.9, 4.0], 8.3)

    doc.add_page_break()
    heading(doc, "Review Actions")
    body(doc, "The print workflow is ready for continued review against the current cycle records. Operational release still depends on approved staffing data and, if required by the examination office, an approved midday session slot. The 2024/2025 timetable should only be reproduced after its source records are provided and explicitly selected.")
    table(doc, ["Priority", "Action", "Owner input required"], [
        ("1", "Confirm chief invigilator assignments for the four scheduled rows.", "Approved staffing assignments linked to event, date, session, and venue."),
        ("2", "Confirm whether 12:00–14:00 is part of the current examination cycle.", "Approved time-slot record or written confirmation that it is not used."),
        ("3", "Confirm the CSC 405 programme assignment.", "Approved course registration or programme record."),
        ("4", "If the photo must be reproduced, provide the approved 2024/2025 source records.", "Approved course, date, venue, population, and staffing records."),
    ], [0.55, 3.45, 2.9], 8.3)
    heading(doc, "Working Tree Note", 2)
    body(doc, "The documentation and QA artifacts were generated without committing, pushing, deploying, resetting the database, logging out, closing the browser, or discarding existing work. The repository already contained other uncommitted changes, which were preserved.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "Branch: main   |   HEAD: 6fb0353   |   Working tree: modified and untracked files present", 8.5, True, NAVY)
    doc.core_properties.title = "BOUESTI ExamPilot Timetable Documentation"
    doc.core_properties.subject = "Print workflow review and QA evidence"
    doc.core_properties.author = "BOUESTI ExamPilot"
    doc.core_properties.keywords = "BOUESTI, ExamPilot, timetable, QA, export"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    build()
